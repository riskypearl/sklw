"""Reads the fetched sheet_workbook.xlsx (see fetch_sheet_workbook.py)
and figures out, for your club's CURRENT week fixture: who's GK, who's
Strikers, on BOTH sides -- then resolves those handles to FPL manager
IDs via clubs.json and writes them straight into matchup_pins.json, the
same file sklw_matchup.py's interactive prompts save to. No manual
lookup on the FPL website, no manual GK/Strikers prompts.

Standalone, pure logic (no network) -- fully testable against a
synthetic workbook built with openpyxl, unlike fetch_sheet_workbook.py
which needs a real login. Separate from that script and from
sklw_matchup.py per this project's usual convention.

How the sheet is read (confirmed against a real example, see README):
  - "Live Scores" tab: one row per fixture, columns (in no particular
    fixed order -- found by scanning for text, not fixed column
    letters, since the exact layout couldn't be inspected ahead of
    time): an "M<N>" label, and the two clubs' names. Whichever row
    contains your club's name tells you which "M<N>" tab is your
    fixture this week, and who the opponent is.
  - Each "M<N>" tab: two clubs stacked, the one named in the TOP banner
    always occupies the LEFT-hand columns, the one in the BOTTOM banner
    the RIGHT-hand columns. Within each club's block: the first 3
    manager rows are GK+Strikers, the next 11 are Squad, the last 2 are
    Bench (16 total) -- confirmed fixed across weeks. GK vs Strikers
    isn't labelled in text at all, only by cell fill color: of the 3
    top rows, the two sharing an IDENTICAL fill color are Strikers, the
    one with a DIFFERENT fill color is GK. Deliberately relative (same
    vs different), not hardcoded to a specific hex color, so a color
    scheme tweak next season doesn't silently break this.

This is automated color/position parsing of someone else's spreadsheet
layout, not an official export format -- ALWAYS sanity-check the first
real run's printed GK/Strikers names against what you can see in the
sheet yourself before trusting it for a decision that matters.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

LIVE_SCORES_SHEET = "Live Scores"


def _text_cells(row_cells) -> list[tuple[int, str]]:
    """(column_index, text) for every cell in a row whose value is a
    non-empty string -- ignores numeric score cells entirely, so we
    never need to know exact column letters ahead of time."""
    out = []
    for cell in row_cells:
        if isinstance(cell.value, str) and cell.value.strip():
            out.append((cell.column, cell.value.strip()))
    return out


def find_fixture(wb, our_club_substring: str) -> tuple[str, str, str]:
    """Scans the Live Scores sheet for the row naming our club (a
    case-insensitive substring match, same style as sklw_matchup.py's
    --opponent). Returns (m_label, our_full_name, opponent_full_name).
    Errors clearly (not a silent guess) if the tab is missing, or if
    our club doesn't appear in exactly one row."""
    if LIVE_SCORES_SHEET not in wb.sheetnames:
        print(f"ERROR: no '{LIVE_SCORES_SHEET}' tab in this workbook -- "
              f"is sheet_workbook.xlsx actually the SKLW master sheet?")
        sys.exit(1)
    ws = wb[LIVE_SCORES_SHEET]
    needle = our_club_substring.strip().casefold()
    matches: list[tuple[str, str, str]] = []
    for row in ws.iter_rows():
        cells = _text_cells(row)
        if len(cells) < 2:
            continue
        m_label = next((t for _, t in cells if t.upper().startswith("M") and t[1:].isdigit()), None)
        names = [t for _, t in cells if t != m_label]
        if m_label is None or len(names) != 2:
            continue
        for i, name in enumerate(names):
            if needle in name.casefold():
                matches.append((m_label, name, names[1 - i]))
    if not matches:
        print(f"ERROR: '{our_club_substring}' didn't match any club in "
              f"the {LIVE_SCORES_SHEET} tab this week.")
        sys.exit(1)
    if len(matches) > 1:
        found = ", ".join(f"{m} ({us} vs {them})" for m, us, them in matches)
        print(f"ERROR: '{our_club_substring}' matched more than one row in "
              f"{LIVE_SCORES_SHEET}, need exactly 1: {found}")
        sys.exit(1)
    return matches[0]


def _classify_top_three(entries: list[tuple[str, str]]) -> tuple[str, list[str]]:
    """entries: [(handle, fill_color), ...] for exactly 3 rows. Returns
    (gk_handle, [striker_handle, striker_handle]) using the "2 alike, 1
    different" fill-color rule -- see module docstring. Errors clearly
    if the colors don't split cleanly into a 2-1 pattern (sheet layout
    assumption broken -- don't guess)."""
    if len(entries) != 3:
        print(f"ERROR: expected exactly 3 people in the GK+Strikers block, "
              f"found {len(entries)}: {[h for h, _ in entries]}")
        sys.exit(1)
    colors = [c for _, c in entries]
    counts = {c: colors.count(c) for c in set(colors)}
    twos = [c for c, n in counts.items() if n == 2]
    ones = [c for c, n in counts.items() if n == 1]
    if len(twos) != 1 or len(ones) != 1:
        print(f"ERROR: couldn't tell GK from Strikers by color -- expected "
              f"2 rows sharing one fill color and 1 with a different "
              f"color, got fill colors {colors} for {[h for h, _ in entries]}. "
              f"Sheet's color scheme may have changed -- check manually.")
        sys.exit(1)
    strikers = [h for h, c in entries if c == twos[0]]
    gk = next(h for h, c in entries if c == ones[0])
    return gk, strikers


def _fill_color(cell) -> str:
    fill = cell.fill
    if fill is None or fill.fgColor is None:
        return ""
    rgb = fill.fgColor.rgb
    return rgb if isinstance(rgb, str) else str(fill.fgColor.theme) + "-" + str(fill.fgColor.tint)


def _club_block(ws, handle_col: int, start_row: int) -> dict[str, list[str] | str]:
    """Reads one club's 16-row block (GK+Strikers, Squad, Bench) from a
    fixed handle column, starting at start_row. Returns
    {"gk": handle, "strikers": [h, h], "squad": [11 handles], "bench":
    [2 handles]}."""
    rows = []
    for r in range(start_row, start_row + 16):
        cell = ws.cell(row=r, column=handle_col)
        if isinstance(cell.value, str) and cell.value.strip():
            rows.append((cell.value.strip(), _fill_color(cell)))
    if len(rows) != 16:
        print(f"ERROR: expected 16 manager rows starting at row {start_row}, "
              f"column {handle_col}, found {len(rows)}. Sheet layout may "
              f"not match what this script expects -- check manually.")
        sys.exit(1)
    gk, strikers = _classify_top_three(rows[:3])
    squad = [h for h, _ in rows[3:14]]
    bench = [h for h, _ in rows[14:16]]
    return {"gk": gk, "strikers": strikers, "squad": squad, "bench": bench}


def parse_matchup_sheet(wb, m_label: str) -> dict:
    """Opens the "M<N>" tab and returns {"left_club": name, "right_club":
    name, "left": {...club block...}, "right": {...club block...}} --
    "left" is whichever club's name is in the TOP banner (see module
    docstring), "right" the one in the bottom banner."""
    if m_label not in wb.sheetnames:
        print(f"ERROR: no '{m_label}' tab in this workbook.")
        sys.exit(1)
    ws = wb[m_label]

    banner_rows = [(r[0].row, r[0].value.strip())
                   for r in ws.iter_rows(max_col=1)
                   if isinstance(r[0].value, str) and r[0].value.strip()]
    # Fallback: banner text might not be in column 1 -- scan every column
    # of every row for a lone long text cell (a club-name banner) if the
    # column-1 scan didn't find at least 2.
    if len(banner_rows) < 2:
        banner_rows = []
        for row in ws.iter_rows():
            cells = _text_cells(row)
            if len(cells) == 1:
                banner_rows.append((row[0].row, cells[0][1]))
    if len(banner_rows) < 2:
        print(f"ERROR: couldn't find two club-name banner rows in '{m_label}' "
              f"-- sheet layout may not match what this script expects.")
        sys.exit(1)
    top_row, left_club = banner_rows[0]
    bottom_row, right_club = banner_rows[-1]

    # Find the handle columns: the first text cell (excluding the
    # banners) below the top banner is the left club's handle column;
    # the last text cell in that same row is the right club's.
    data_row = None
    left_col = right_col = None
    for row in ws.iter_rows(min_row=top_row + 1, max_row=bottom_row - 1):
        cells = _text_cells(row)
        if len(cells) >= 2:
            data_row = row[0].row
            left_col = cells[0][0]
            right_col = cells[-1][0]
            break
    if data_row is None:
        print(f"ERROR: couldn't find any manager rows in '{m_label}'.")
        sys.exit(1)

    left = _club_block(ws, left_col, data_row)
    right = _club_block(ws, right_col, data_row)
    return {"left_club": left_club, "right_club": right_club, "left": left, "right": right}


def resolve_ids(handles: list[str], roster: dict[str, int], label: str) -> list[int]:
    """Matches sheet handles to clubs.json's roster keys case-
    insensitively (sheet formatting can differ slightly in case from
    the master list) -- warns and drops any handle that doesn't match
    rather than silently guessing."""
    by_fold = {k.casefold(): v for k, v in roster.items()}
    ids = []
    for h in handles:
        mid = by_fold.get(h.casefold())
        if mid is None:
            print(f"WARNING: '{h}' ({label}) not found in this club's clubs.json "
                  f"roster -- skipped. Rebuild clubs.json if the master list "
                  f"has changed, or check for a name mismatch.")
            continue
        ids.append(mid)
    return ids


def load_pins(path: str) -> dict:
    p = Path(path)
    if not p.exists():
        return {"us": {}, "them": {}}
    try:
        data = json.loads(p.read_text())
    except json.JSONDecodeError:
        return {"us": {}, "them": {}}
    data.setdefault("us", {})
    data.setdefault("them", {})
    return data


def save_pins(path: str, pins: dict) -> None:
    Path(path).write_text(json.dumps(pins, indent=2))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default="sheet_workbook.xlsx",
                     help="path to the xlsx fetched by fetch_sheet_workbook.py")
    ap.add_argument("--clubs-file", default="clubs.json")
    ap.add_argument("--our-club", required=True,
                     help="substring of our club's name, e.g. 'Algorithm'")
    ap.add_argument("--pins-file", default="matchup_pins.json")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl isn't installed -- run: pip install openpyxl")
        sys.exit(1)

    if not Path(args.workbook).exists():
        print(f"ERROR: no workbook at {args.workbook} -- run "
              f"fetch_sheet_workbook.py first.")
        sys.exit(1)
    if not Path(args.clubs_file).exists():
        print(f"ERROR: no clubs.json at {args.clubs_file} -- run "
              f"build_clubs_json.py first.")
        sys.exit(1)

    clubs = json.loads(Path(args.clubs_file).read_text())
    wb = openpyxl.load_workbook(args.workbook, data_only=True)

    m_label, our_name, opponent_name = find_fixture(wb, args.our_club)
    print(f"This week's fixture: {our_name} vs {opponent_name} ({m_label})")

    sheet = parse_matchup_sheet(wb, m_label)
    if args.our_club.strip().casefold() in sheet["left_club"].casefold():
        us_block, them_block = sheet["left"], sheet["right"]
        us_club_name, them_club_name = sheet["left_club"], sheet["right_club"]
    else:
        us_block, them_block = sheet["right"], sheet["left"]
        us_club_name, them_club_name = sheet["right_club"], sheet["left_club"]

    def club_roster(name: str) -> dict[str, int]:
        found = [c for c in clubs if name.casefold() in c.casefold() or c.casefold() in name.casefold()]
        if len(found) != 1:
            print(f"ERROR: '{name}' matched {len(found)} club(s) in clubs.json, "
                  f"need exactly 1: {found}")
            sys.exit(1)
        return clubs[found[0]]

    us_roster = club_roster(us_club_name)
    them_roster = club_roster(them_club_name)

    us_gk_ids = resolve_ids([us_block["gk"]], us_roster, "our GK")
    us_striker_ids = resolve_ids(us_block["strikers"], us_roster, "our Strikers")
    them_gk_ids = resolve_ids([them_block["gk"]], them_roster, "their GK")
    them_striker_ids = resolve_ids(them_block["strikers"], them_roster, "their Strikers")

    if not us_gk_ids or len(us_striker_ids) != 2 or not them_gk_ids or len(them_striker_ids) != 2:
        print("ERROR: couldn't resolve every GK/Strikers handle to an ID "
              "(see warnings above) -- nothing written to matchup_pins.json.")
        sys.exit(1)

    print(f"Us  ({us_club_name}): GK={us_block['gk']}, Strikers={us_block['strikers']}")
    print(f"Them ({them_club_name}): GK={them_block['gk']}, Strikers={them_block['strikers']}")

    pins = load_pins(args.pins_file)
    pins["us"] = {"gk_id": str(us_gk_ids[0]),
                  "strikers_ids": ",".join(str(i) for i in us_striker_ids)}
    pins["them"][them_club_name] = {"gk_id": str(them_gk_ids[0]),
                                     "strikers_ids": ",".join(str(i) for i in them_striker_ids)}
    save_pins(args.pins_file, pins)
    print(f"\nSaved to {args.pins_file}. Run: "
          f"python sklw_matchup.py --opponent \"{them_club_name}\"")


if __name__ == "__main__":
    main()
